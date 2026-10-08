# -*- coding: utf-8 -*-
"""05 — Alinha a coluna Title da planilha master (Drive) com os títulos do site.

O site é a referência editorial (títulos padronizados em 01/08/2026:
"Peça - Marca SKU | Aplicação" / kits "Peça - Kit N Unidades - Marca SKU").
A planilha segue o site; preço continua sendo domínio da planilha.

Também corrige o SKU 5.0220.0548824.0 → 5.0220.0548824-2 (SKU do site nunca muda).

Uso:
    python3 05_alinhar_titulos_planilha.py           # audita (não grava)
    python3 05_alinhar_titulos_planilha.py --local   # audita usando integra/central.xlsx em cache
    python3 05_alinhar_titulos_planilha.py --push    # grava no Drive

Requer token do admin@ com escopo Drive:
    gcloud auth login admin@agropecaspadrao.com.br --enable-gdrive-access --force
    (sem --force o gcloud reusa a credencial antiga sem o escopo)

ATENÇÃO (incidente 01/08/2026): salvar xlsx com openpyxl APAGA o cache de valores
das fórmulas — leituras data_only passam a ver None. Depois de --push, rode o
reparo: Drive files.copy convertendo p/ Google Sheets → export xlsx → PATCH de
volta no mesmo file id → deletar a cópia (recalcula e restaura o cache).
"""
import sys, os, re, io, urllib.request
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'integra'))
import pipeline as P
import openpyxl

PUSH = '--push' in sys.argv
SKU_FIX = {'5.0220.0548824.0': '5.0220.0548824-2'}

# ---- site: título por (sku_base, kit) ----
prods = P.pull_shopify()
site = {}
for p in prods:
    if p['status'] != 'ACTIVE': continue
    for v in p['variants']['nodes']:
        sku = (v['sku'] or '').strip()
        if not sku: continue
        m = re.search(r'-KIT(\d+)$', sku, re.I)
        base = sku[:m.start()] if m else sku
        kit = int(m.group(1)) if m else 1
        site[(base, kit)] = p['title'].strip()
print(f'{len(site)} (SKU, kit) ativos no site')

# ---- baixa a master preservando fórmulas ----
if '--local' in sys.argv:
    raw = open(os.path.join(os.path.dirname(P.__file__), 'central.xlsx'), 'rb').read()
    tk = None
    if PUSH: sys.exit('--local não combina com --push')
else:
    tk = P.gtoken()
    url = f'https://www.googleapis.com/drive/v3/files/{P.SHEET_ID}?alt=media&supportsAllDrives=true'
    raw = urllib.request.urlopen(urllib.request.Request(url, headers={'Authorization': 'Bearer ' + tk})).read()
wb = openpyxl.load_workbook(io.BytesIO(raw))  # sem data_only: preserva fórmulas

def site_hit(sku, kit):
    for s in P.sku_variants(sku):
        if (s, kit) in site: return site[(s, kit)]
    # mapeia .0 ↔ -2 (bomba dupla)
    alt = SKU_FIX.get(sku)
    if alt and (alt, kit) in site: return site[(alt, kit)]
    return None

n_tit = n_sku = 0
for name in P.SUP:
    if name not in wb.sheetnames: continue
    ws = wb[name]
    hr = next((r for r in range(1, 9)
               if any('Title' in str(ws.cell(r, c).value or '') for c in range(1, ws.max_column + 1))), 5)
    cols = {str(ws.cell(hr, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(hr, c).value}
    cTit = next((c for k, c in cols.items() if 'Title' in k), None)
    cKit = next((c for k, c in cols.items() if 'Qtde KIT' in k), None)
    if not cTit: continue
    for r in range(hr + 1, ws.max_row + 1):
        sku = ws.cell(r, 1).value
        if sku in (None, ''): continue
        sku = str(sku).strip()
        kit = ws.cell(r, cKit).value if cKit else 1
        try: kit = int(float(kit)) if kit else 1
        except (ValueError, TypeError): kit = 1
        novo = site_hit(sku, kit)
        if novo and str(ws.cell(r, cTit).value or '').strip() != novo:
            print(f'  {name} L{r} [{sku} kit{kit}] Title → {novo}')
            if PUSH: ws.cell(r, cTit).value = novo
            n_tit += 1
        if sku in SKU_FIX:
            print(f'  {name} L{r} SKU {sku} → {SKU_FIX[sku]}')
            if PUSH: ws.cell(r, 1).value = SKU_FIX[sku]
            n_sku += 1

print(f"\n{'GRAVANDO' if PUSH else 'AUDIT (nada gravado)'}: {n_tit} títulos, {n_sku} SKUs")
if PUSH and (n_tit or n_sku):
    buf = io.BytesIO(); wb.save(buf)
    req = urllib.request.Request(
        f'https://www.googleapis.com/upload/drive/v3/files/{P.SHEET_ID}?uploadType=media&supportsAllDrives=true',
        data=buf.getvalue(), method='PATCH',
        headers={'Authorization': 'Bearer ' + tk,
                 'Content-Type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'})
    urllib.request.urlopen(req)
    print('Planilha master atualizada no Drive.')
