"""Reconciliador do ml_publish_report.csv -> classificação + CSV corrigido.

Classifica cada SKU (regex em code/cause, tolera JSON truncado), separa o que é
auto-corrigível do que é humano, consulta a API do ML (category predictor +
diagnóstico de conta) e gera o CSV pronto para reenvio (OK intactos).

Uso: python3 reconcile_report.py
"""
from __future__ import annotations

import csv
import json
import re

import requests

import config
from ml_auth import get_access_token

API = "https://api.mercadolibre.com"
BRAND = "APP — Agro Peças Padrão"          # regra da APP: marca sempre esta
MIN_PRICE = 8.0

# ---------- fontes de dados ----------
site = {p["sku"]: p for p in json.loads(config.PRODUCTS_JSON.read_text(encoding="utf-8"))}
report = list(csv.DictReader(config.PUBLISH_REPORT.open(encoding="utf-8")))


def norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


# preço mestre (Preço final site) por SKU normalizado
def load_master_price() -> dict[str, float]:
    from openpyxl import load_workbook
    wb = load_workbook(config.ROOT / "util" / "APP_Master_Produtos_Shopify.xlsx", data_only=True)
    out: dict[str, float] = {}
    for name in [n for n in wb.sheetnames if n[:2] in ("02", "03", "04", "05", "06", "07")]:
        ws = wb[name]
        hrow = next((ri for ri in range(1, 8)
                     if any("SKU / Código" == str(ws.cell(row=ri, column=c).value)
                            for c in range(1, ws.max_column + 1))), None)
        if not hrow:
            continue
        hdr = {c: ws.cell(row=hrow, column=c).value for c in range(1, ws.max_column + 1)}
        c_sku = next((c for c, h in hdr.items() if h == "SKU / Código"), None)
        c_price = next((c for c, h in hdr.items() if isinstance(h, str) and "final site" in h.lower()), None)
        if not (c_sku and c_price):
            continue
        for ri in range(hrow + 1, ws.max_row + 1):
            sku = ws.cell(row=ri, column=c_sku).value
            val = ws.cell(row=ri, column=c_price).value
            if sku and isinstance(val, (int, float)):
                out[norm(sku)] = float(val)
    return out


# ---------- extração via regex (JSON pode estar truncado) ----------
def extract(motivo: str) -> dict:
    codes = re.findall(r'"code"\s*:\s*"([^"]+)"', motivo)
    miss = re.search(r'attributes \[([^\]]+)\] are required for category (MLB\d+)', motivo)
    price_cat = re.search(r'category (MLB\d+) requires a minimum of price', motivo)
    cat = re.search(r'cat=(MLB\d+)', motivo) or re.search(r'category (MLB\d+)', motivo)
    return {
        "codes": codes,
        "address_pending": ("address_pending" in motivo or "unable_to_list" in motivo),
        "missing_attr": miss.group(1) if miss else None,
        "category_id": (cat.group(1) if cat else None),
        "price_floor_cat": (price_cat.group(1) if price_cat else None),
        "no_category": ("categoria não prevista" in motivo),
    }


def classify(r: dict, ex: dict) -> str:
    if r["status"] == "ok":
        return "PUBLICADO"
    if r["status"] == "skip" and "sem foto" in (r["motivo"] or ""):
        return "FOTO"
    has_info = ex["no_category"] or any(
        c in ("item.attributes.missing_required", "item.price.invalid") for c in ex["codes"])
    has_conta = ex["address_pending"] or any(c.startswith("shipping") for c in ex["codes"])
    if has_info and has_conta:
        return "CONTA+INFO"
    if has_info:
        return "INFO"
    if has_conta:
        return "CONTA_FRETE"
    return "OUTRO"


# ---------- ML API ----------
def ml_predict_category(token: str, title: str) -> str | None:
    base = title.split("—")[0]
    clean = " ".join(w for w in base.split() if not any(c.isdigit() for c in w))[:60]
    for q in dict.fromkeys([title[:60], clean]):
        if not q.strip():
            continue
        try:
            r = requests.get(f"{API}/sites/{config.ML_SITE}/domain_discovery/search",
                             params={"q": q, "limit": 1},
                             headers={"Authorization": f"Bearer {token}"}, timeout=20)
            if r.status_code == 200 and r.json():
                return r.json()[0].get("category_id")
        except requests.RequestException:
            pass
    return None


def ml_account_diagnostics(token: str) -> dict:
    try:
        r = requests.get(f"{API}/users/me", headers={"Authorization": f"Bearer {token}"}, timeout=20)
        u = r.json() if r.status_code == 200 else {}
    except requests.RequestException:
        u = {}
    return {
        "nickname": u.get("nickname"),
        "status_list": (u.get("status") or {}).get("list", {}),
        "status_site": (u.get("status") or {}).get("site_status"),
        "address": u.get("address"),
        "shipping_modes": (u.get("seller_reputation") or {}).get("metrics") is not None,
        "raw_status": u.get("status"),
    }


def main() -> None:
    token = get_access_token()
    master = load_master_price()
    acct = ml_account_diagnostics(token)

    recs = []
    for r in report:
        sku = r["sku"]
        ex = extract(r["motivo"] or "")
        cat_prim = classify(r, ex)
        p = site.get(sku, {})
        part_number = p.get("oem") or sku
        price_pub = p.get("price")
        price_master = master.get(norm(sku))
        cat_id = ex["category_id"]

        # ---- decisão de ação / responsável / auto ----
        acao, resp, auto, resubmeter, bloqueio, price_sug = "", "", "", "", "", price_pub

        if cat_prim == "PUBLICADO":
            acao, resp, resubmeter = "Já no ar — manter intacto", "—", "nao"
        elif cat_prim == "FOTO":
            acao = f"Subir ≥1 foto no Shopify (handle: {p.get('handle','?')})"
            resp, auto, resubmeter, bloqueio = "HUMANO (fotografia)", "nao", "nao", "foto"
        else:  # INFO, CONTA_FRETE, CONTA+INFO
            partes = []
            # categoria
            if ex["no_category"] and not cat_id:
                cat_id = ml_predict_category(token, p.get("title") or sku)
                partes.append("categoria via predictor ML" + (f" -> {cat_id}" if cat_id else " (não resolveu: manual)"))
            # atributo faltante
            if ex["missing_attr"] == "PART_NUMBER":
                partes.append(f"PART_NUMBER={part_number}")
            elif ex["missing_attr"]:
                partes.append(f"atributo {ex['missing_attr']} (valor manual)")
            # preço
            if "item.price.invalid" in ex["codes"]:
                if price_master and price_master >= MIN_PRICE:
                    price_sug = price_master
                    partes.append(f"preço {price_pub}->{price_master} (planilha mestre)")
                else:
                    partes.append(f"preço {price_pub} < R${MIN_PRICE:.0f}: REVISAR (subir/kit)")
            acao = "; ".join(partes) or "rever payload"
            # auto-fixável se todas as partes forem automáticas
            human_bits = ("manual" in acao) or ("REVISAR" in acao)
            auto = "nao" if human_bits else "sim"
            resp = "AGENTE" if auto == "sim" else "AGENTE+HUMANO"
            if cat_prim == "CONTA_FRETE":
                resp, auto, bloqueio = "CONTA ML", "nao", "conta"
                acao = "(sem ação por SKU — ver ajustes de conta)"
                resubmeter = "após conta"
            elif cat_prim == "CONTA+INFO":
                bloqueio, resubmeter = "conta", "após conta"
            else:  # INFO puro
                resubmeter = "sim"

        recs.append({
            "sku": sku, "categoria_primaria": cat_prim, "resubmeter": resubmeter,
            "bloqueio": bloqueio, "ml_id": r["ml_id"], "brand": BRAND,
            "ml_title": (p.get("title") or "")[:60],
            "price_publicado": price_pub, "price_master": price_master, "price_sugerido": price_sug,
            "part_number": part_number, "ml_category_id": cat_id or "",
            "atributo_faltante": ex["missing_attr"] or "", "n_fotos": p.get("n_images", 0),
            "codes": "|".join(ex["codes"]), "acao": acao, "responsavel": resp, "auto_fixavel": auto,
        })

    # ---- escreve CSV corrigido ----
    fields = ["sku", "categoria_primaria", "resubmeter", "bloqueio", "ml_id", "brand", "ml_title",
              "price_publicado", "price_master", "price_sugerido", "part_number", "ml_category_id",
              "atributo_faltante", "n_fotos", "codes", "acao", "responsavel", "auto_fixavel"]
    out = config.DATA_DIR / "ml_produtos_corrigido.csv"
    with out.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(recs)

    # ---- resumo ----
    from collections import Counter
    tot = len(recs)
    c = Counter(x["categoria_primaria"] for x in recs)
    print("=" * 56)
    print(f"RESUMO — {tot} SKUs")
    print("=" * 56)
    for k in ["PUBLICADO", "FOTO", "INFO", "CONTA_FRETE", "CONTA+INFO", "OUTRO"]:
        if c.get(k):
            print(f"  {c[k]:3d}  {k:<12} {c[k]*100/tot:5.1f}%")
    pub = c.get("PUBLICADO", 0)
    print(f"\n  Taxa de sucesso (publicado): {pub*100//tot}%   Taxa de falha: {(tot-pub)*100//tot}%")
    auto = sum(1 for x in recs if x["auto_fixavel"] == "sim")
    print(f"  Auto-corrigíveis pelo agente: {auto}")
    print(f"\n  Diagnóstico conta ML: nickname={acct['nickname']} site_status={acct['status_site']}")
    print(f"    status.list: {acct['status_list']}")
    print(f"\nCSV corrigido -> {out}")


if __name__ == "__main__":
    main()
